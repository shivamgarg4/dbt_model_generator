
from openpyxl import load_workbook
import snowflake.connector
import os



def generate_lnd_dbt_model_file(config, file_path):
    """Generate the dbt model file."""
    try:
        # Extract source information
        source_name = 'source'

        workbook = load_workbook(file_path)
        pipeline_sheet = workbook.active  # Assuming the pipeline details are in the first sheet

        # Find the header row and column indices
        header_row = 1
        source_model_col = None
        target_model_col = None
        model_type_col = None

        for col in range(1, pipeline_sheet.max_column + 1):
            header_value = pipeline_sheet.cell(row=header_row, column=col).value
            if header_value == 'Source Model':
                source_model_col = col
            elif header_value == 'Target Model':
                target_model_col = col
            elif header_value == 'Model Type':
                model_type_col = col

        if not all([source_model_col, target_model_col, model_type_col]):
            raise ValueError("Required columns not found in the sheet")

        # Get source and target information from the first LND Load type row
        source_table = None
        target_table = None

        for row in range(2, pipeline_sheet.max_row + 1):
            model_type = pipeline_sheet.cell(row=row, column=model_type_col).value
            if model_type == 'LND Load':
                source_table = pipeline_sheet.cell(row=row, column=source_model_col).value
                target_table = pipeline_sheet.cell(row=row, column=target_model_col).value
                break

        if not source_table or not target_table:
            raise ValueError("Source or target table not found for LND Load type")

        # Process source and target information
        source_parts = source_table.split('.')
        source_db, source_schema, source_table_name = source_parts[0], source_parts[1], source_parts[2]

        target_parts = target_table.split('.')
        target_schema = target_parts[0] if len(target_parts) > 1 else None
        target_table_name = target_parts[-1]

        # The rest of the code remains the same from here
        """Extract Snowflake configuration from Config sheet"""
        config_sheet = workbook['Config']  # Assuming Config sheet still exists
        snowflake_config = {}
        in_snowflake_section = False

        for row in range(2, config_sheet.max_row + 1):
            param = config_sheet.cell(row=row, column=1).value
            if not param:
                continue
            if param == 'Snowflake Configuration':
                in_snowflake_section = True
                continue
            if in_snowflake_section and param in ['ROLE', 'WAREHOUSE', 'DATABASE', 'ACCOUNT', 'USER', 'AUTHENTICATOR']:
                value = config_sheet.cell(row=row, column=2).value
                if value:
                    snowflake_config[param] = value

        # Connect to Snowflake and get column information
        conn = snowflake.connector.connect(
            account=snowflake_config['ACCOUNT'],
            user=snowflake_config['USER'],
            authenticator=snowflake_config['AUTHENTICATOR'],
            warehouse=snowflake_config['WAREHOUSE'],
            role=snowflake_config['ROLE'],
            database=source_db,
            schema=source_schema
        )

        cursor = conn.cursor()
        cursor.execute(f"""
            SELECT COLUMN_NAME
            FROM {source_db}.INFORMATION_SCHEMA.COLUMNS
            WHERE TABLE_SCHEMA = '{source_schema}'
            AND TABLE_NAME = '{source_table_name}'
            ORDER BY ORDINAL_POSITION
        """)

        columns = cursor.fetchall()
        formatted_columns = format_columns(columns)
        replaced_columns = replace_audit_columns(formatted_columns, source_schema, source_table_name, target_schema,
                                                 target_table_name)

        model_config = f"""
        {{
            {{
                config(
                    schema = '{target_schema}',
                    tags=['{target_table_name}'],
                    alias='{target_table_name}',
                    materialized='table',
                    transient = false
                )
            }}
        }}
        -- Model: {target_schema}.{target_table_name}
        -- Source: {source_db}.{source_schema}.{source_table_name}
        """

        model_config += "\n\nSELECT\n"
        model_config += ",\n".join(replaced_columns) + "\n"
        model_config += f"FROM {{{{ source('{source_name}', '{source_table_name}') }}}}\n"

        # Create output directory if it doesn't exist
        output_dir = 'models'
        os.makedirs(output_dir, exist_ok=True)

        # Generate file name
        model_name = f"{target_schema}.{target_table_name}"
        file_name = f"{model_name}.sql"
        file_path = os.path.join(output_dir, file_name)

        # Write model file
        with open(file_path, 'w') as f:
            f.write(model_config)

        return True, file_path
    except Exception as e:
        raise Exception(f"An error occurred while generating the dbt model: {e}")


def format_columns(columns):
    """Format columns for the dbt model."""
    try:
        max_length = max(len(column[0]) for column in columns)
        return [f"    {column[0].ljust(max_length)}" for column in columns]
    except Exception as e:
        raise Exception(f"An error occurred while formatting columns: {e}")


def replace_audit_columns(columns, source_schema_name, source_view_name, target_schema_name, target_table_name):
    """Replace audit columns with static values."""
    try:
        audit_columns = {
            "DATA_SRC": f"'{source_schema_name}.{source_view_name}' AS DATA_SRC",
            "CREATE_DT": "CURRENT_TIMESTAMP() AS CREATE_DT",
            "CREATE_BY": "CAST(CURRENT_USER() AS VARCHAR(200)) AS CREATE_BY",
            "CREATE_PGM": f"'{target_schema_name}.{target_table_name}' AS CREATE_PGM",
            "UPDATE_DT": "CURRENT_TIMESTAMP() AS UPDATE_DT",
            "UPDATE_BY": "CAST(CURRENT_USER() AS VARCHAR(200)) AS UPDATE_BY",
            "UPDATE_PGM": f"'{target_schema_name}.{target_table_name}' AS UPDATE_PGM"
        }
        return [f"    {audit_columns.get(column.split()[0].strip(), column)}" for column in columns]
    except Exception as e:
        raise Exception(f"An error occurred while replacing audit columns: {e}")
